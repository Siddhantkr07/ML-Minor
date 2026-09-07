"""Build the final submission report from a completed experiment's artifacts."""

from __future__ import annotations

import json
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


PROJECT_ROOT = Path(__file__).resolve().parents[1]
ARTIFACTS = PROJECT_ROOT / "artifacts"
REPORTS = PROJECT_ROOT / "reports"
REPORT_PATH = REPORTS / "Stock_Price_Prediction_Report.docx"

NAVY = "1F4E78"
PALE_BLUE = "EAF2F8"
LIGHT_GRAY = "D9D9D9"
BLACK = RGBColor(0, 0, 0)


def set_cell_shading(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_border(cell, color: str = LIGHT_GRAY) -> None:
    properties = cell._tc.get_or_add_tcPr()
    borders = properties.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        properties.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "6")
        element.set(qn("w:color"), color)


def set_cell_padding(cell, amount: int = 110) -> None:
    properties = cell._tc.get_or_add_tcPr()
    margins = properties.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        properties.append(margins)
    for side in ("top", "start", "bottom", "end"):
        element = margins.find(qn(f"w:{side}"))
        if element is None:
            element = OxmlElement(f"w:{side}")
            margins.append(element)
        element.set(qn("w:w"), str(amount))
        element.set(qn("w:type"), "dxa")


def set_run_font(run, size: float | None = None, bold: bool | None = None, color: RGBColor = BLACK) -> None:
    run.font.name = "Aptos"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    run.font.color.rgb = color
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold


def add_paragraph(document: Document, text: str = "", style: str | None = None, alignment=None):
    paragraph = document.add_paragraph(style=style)
    if alignment is not None:
        paragraph.alignment = alignment
    if text:
        set_run_font(paragraph.add_run(text), size=10.5)
    return paragraph


def add_bullet(document: Document, text: str) -> None:
    paragraph = document.add_paragraph(style="List Bullet")
    paragraph.paragraph_format.space_after = Pt(3)
    set_run_font(paragraph.add_run(text), size=10.5)


def add_heading(document: Document, text: str, level: int = 1) -> None:
    paragraph = document.add_paragraph(style=f"Heading {level}")
    paragraph.paragraph_format.keep_with_next = True
    set_run_font(paragraph.add_run(text), size=15 if level == 1 else 12.5, bold=True)


def add_table(document: Document, headers: list[str], rows: list[list[str]], widths: list[float] | None = None) -> None:
    table = document.add_table(rows=1, cols=len(headers))
    table.autofit = False
    table.style = "Table Grid"
    header_cells = table.rows[0].cells
    for index, header in enumerate(headers):
        cell = header_cells[index]
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        set_cell_shading(cell, NAVY)
        set_cell_border(cell)
        set_cell_padding(cell)
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_run_font(paragraph.add_run(header), size=9.5, bold=True, color=RGBColor(255, 255, 255))
        if widths:
            cell.width = Inches(widths[index])

    for row_index, row in enumerate(rows):
        cells = table.add_row().cells
        for column_index, value in enumerate(row):
            cell = cells[column_index]
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_border(cell)
            set_cell_padding(cell)
            if row_index % 2 == 1:
                set_cell_shading(cell, PALE_BLUE)
            paragraph = cell.paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if column_index > 0 else WD_ALIGN_PARAGRAPH.LEFT
            set_run_font(paragraph.add_run(value), size=9.5)
            if widths:
                cell.width = Inches(widths[column_index])
    document.add_paragraph().paragraph_format.space_after = Pt(3)


def configure_styles(document: Document) -> None:
    normal = document.styles["Normal"]
    normal.font.name = "Aptos"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = BLACK
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.15

    for name, size in (("Title", 23), ("Heading 1", 15), ("Heading 2", 12.5)):
        style = document.styles[name]
        style.font.name = "Aptos Display" if name == "Title" else "Aptos"
        style._element.rPr.rFonts.set(qn("w:ascii"), style.font.name)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), style.font.name)
        style.font.size = Pt(size)
        style.font.color.rgb = BLACK
        style.font.bold = True

    if "Caption" not in document.styles:
        document.styles.add_style("Caption", WD_STYLE_TYPE.PARAGRAPH)


def add_cover_page(document: Document, source: str, first_date: str, last_date: str) -> None:
    for _ in range(5):
        document.add_paragraph()
    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(title.add_run("Stock Price Prediction Minor Project Report"), size=23, bold=True)

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_before = Pt(12)
    set_run_font(subtitle.add_run("Chronological Evaluation with Ridge Regression"), size=14)

    document.add_paragraph()
    details = [
        "Prepared by Siddhant",
        "Minor Project",
        f"Dataset {source} daily closing prices",
        f"Verified period {first_date} to {last_date}",
        "Repository https://github.com/Siddhantkr07/ML-Minor",
    ]
    for detail in details:
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_run_font(paragraph.add_run(detail), size=11)
    document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def main() -> None:
    metrics_path = ARTIFACTS / "metrics.json"
    summary_path = ARTIFACTS / "dataset_summary.json"
    chart_path = ARTIFACTS / "forecast.png"
    for required_path in (metrics_path, summary_path, chart_path):
        if not required_path.is_file():
            raise FileNotFoundError(f"Required experiment output is missing: {required_path}")

    metrics_payload = json.loads(metrics_path.read_text(encoding="utf-8"))
    summary = json.loads(summary_path.read_text(encoding="utf-8"))
    ridge = metrics_payload["metrics"]["ridge_regression"]
    baseline = metrics_payload["metrics"]["naive_persistence"]
    source = summary["source"]
    first_date = summary["date_range"]["first_observation"]
    last_date = summary["date_range"]["last_observation"]

    REPORTS.mkdir(parents=True, exist_ok=True)
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)
    configure_styles(document)
    document.core_properties.title = "Stock Price Prediction Minor Project Report"
    document.core_properties.author = "Siddhant"
    document.core_properties.subject = "Time-series regression project"

    add_cover_page(document, source, first_date, last_date)

    add_heading(document, "Abstract")
    add_paragraph(
        document,
        "This project evaluates whether a regression model can forecast the next trading day's closing price from recent price history. "
        f"Daily {source} closing-price data from {first_date} to {last_date} was retrieved through Yahoo Finance using yfinance. "
        "The experiment builds lag, return, and rolling-window features, then reserves the newest 20 percent of eligible observations as a chronological holdout set. "
        f"Ridge Regression achieved a mean absolute error of {ridge['mae']:.4f}, root mean squared error of {ridge['rmse']:.4f}, and R-squared of {ridge['r2']:.4f}. "
        f"The naive persistence baseline achieved a lower mean absolute error of {baseline['mae']:.4f}, showing that a simple next-close assumption remained the stronger benchmark for this run. "
        "The result is reported as an educational forecasting experiment and not as investment advice.",
    )

    add_heading(document, "Problem Statement")
    add_paragraph(
        document,
        "The objective is to predict the closing price one trading day ahead using information available at the end of the current trading day. "
        "The project treats this as a time-series regression task and evaluates the model only on data that occurs after the training period.",
    )

    add_heading(document, "Project Objectives")
    for objective in (
        "Retrieve and validate historical daily closing-price data.",
        "Create time-aware features without using future prices as inputs.",
        "Train a reproducible Ridge Regression pipeline.",
        "Evaluate the model against a chronological holdout set and a naive benchmark.",
        "Produce a chart, metrics, saved model, and reusable project documentation.",
    ):
        add_bullet(document, objective)

    add_heading(document, "Dataset")
    add_paragraph(
        document,
        "The verified run uses Apple Inc. daily historical price data with ticker AAPL. The data loader accepts either Yahoo Finance downloads or a local CSV with Date and Close columns. "
        "For the verified experiment, the close series was sorted chronologically, duplicate dates were removed while retaining the final occurrence, non-numeric values were discarded, and positive close values were required.",
    )
    add_table(
        document,
        ["Attribute", "Verified value"],
        [
            ["Data source", "Yahoo Finance through yfinance"],
            ["Ticker", source],
            ["Observed date range", f"{first_date} to {last_date}"],
            ["Trading-day observations", str(summary["observations"])],
            ["Close price range", f"{summary['close_price']['minimum']:.2f} to {summary['close_price']['maximum']:.2f}"],
            ["Mean close price", f"{summary['close_price']['mean']:.2f}"],
        ],
        widths=[2.5, 3.8],
    )

    add_heading(document, "Feature Engineering")
    add_paragraph(
        document,
        "Each model row uses values available at the end of that row's date. The target is shifted one trading day forward, so the model learns to estimate the following close. "
        "Rows without enough history for every required feature are excluded from model fitting.",
    )
    add_table(
        document,
        ["Feature group", "Variables used"],
        [
            ["Current price", "Current closing price"],
            ["Lagged price", "1, 2, 3, 5, 10, and 20 trading-day lags"],
            ["Momentum", "Daily percentage return"],
            ["Rolling statistics", "5, 10, and 20-day mean and standard deviation"],
            ["Relative price", "Distance from each rolling mean"],
        ],
        widths=[2.0, 4.3],
    )

    add_heading(document, "Model and Evaluation Design")
    add_paragraph(
        document,
        "The fitted pipeline performs median imputation, standard scaling, and Ridge Regression with alpha equal to 1.0. "
        "The newest 20 percent of complete rows is held out as the test set. A chronological split is essential for time series because training on later observations would make the forecast evaluation unrealistically optimistic. "
        "The baseline predicts that the next closing price will equal the current closing price.",
    )
    add_table(
        document,
        ["Experiment setting", "Verified value"],
        [
            ["Forecast horizon", f"{metrics_payload['forecast_horizon_trading_days']} trading day"],
            ["Training rows", str(metrics_payload["train_rows"])],
            ["Test rows", str(metrics_payload["test_rows"])],
            ["Primary model", "Median imputation + StandardScaler + Ridge Regression"],
            ["Comparison model", "Naive persistence baseline"],
        ],
        widths=[2.35, 3.95],
    )

    add_heading(document, "Results")
    add_paragraph(
        document,
        "The table below reports out-of-sample performance on the final chronological holdout period. Lower MAE and RMSE are better; higher R-squared is better.",
    )
    add_table(
        document,
        ["Model", "MAE", "RMSE", "R-squared"],
        [
            ["Ridge Regression", f"{ridge['mae']:.4f}", f"{ridge['rmse']:.4f}", f"{ridge['r2']:.4f}"],
            ["Naive persistence", f"{baseline['mae']:.4f}", f"{baseline['rmse']:.4f}", f"{baseline['r2']:.4f}"],
        ],
        widths=[2.5, 1.15, 1.25, 1.4],
    )
    figure = document.add_paragraph()
    figure.alignment = WD_ALIGN_PARAGRAPH.CENTER
    figure.add_run().add_picture(str(chart_path), width=Inches(6.3))
    caption = document.add_paragraph()
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(caption.add_run("Figure 1. Chronological holdout predictions for AAPL."), size=9, bold=True)

    add_heading(document, "Discussion")
    add_paragraph(
        document,
        "The Ridge model follows the broad movement of the actual closing-price series and achieves a high R-squared on this holdout period. "
        "However, the naive persistence baseline has lower MAE and RMSE and a slightly higher R-squared. This means the more complex model did not improve on simply using the current close as the next-close prediction for this particular experiment. "
        "Reporting this comparison is important because it prevents an overstatement of forecasting capability.",
    )

    add_heading(document, "Limitations")
    for limitation in (
        "Only historical closing prices are used; the model has no news, earnings, fundamentals, macroeconomic, or intraday inputs.",
        "The result is based on one asset and one holdout period, so it may not generalize to other stocks or market conditions.",
        "The evaluation does not include trading costs, risk management, or portfolio performance.",
        "Forecasts are uncertain and must not be used as investment advice or an automated trading signal.",
    ):
        add_bullet(document, limitation)

    add_heading(document, "Reproducibility")
    add_paragraph(
        document,
        "The source code, tests, data documentation, model card, and report are available in the project repository. The following command reproduces the verified experiment from the project root after dependencies are installed:",
    )
    command = document.add_paragraph()
    command.paragraph_format.left_indent = Inches(0.25)
    set_run_font(
        command.add_run("$env:PYTHONPATH = \"src\"\n.\\.venv\\Scripts\\python.exe -m stock_predictor.cli train --ticker AAPL --start 2020-01-01 --end 2025-01-01"),
        size=9.5,
    )
    add_paragraph(
        document,
        "Each run writes the cleaned price history, data summary, holdout predictions, chart, metrics, and trained model to the local artifacts directory.",
    )

    add_heading(document, "Conclusion")
    add_paragraph(
        document,
        "This project delivers a reproducible next-day stock-price forecasting workflow with data validation, leakage-aware feature construction, chronological evaluation, automated tests, and documented results. "
        "The verified AAPL run shows that Ridge Regression can closely track price movement, while the naive baseline remains stronger on the selected holdout metrics. Future work should use walk-forward validation, multiple assets, and carefully validated additional predictors.",
    )

    add_heading(document, "References")
    references = (
        "Yahoo Finance data accessed through yfinance. https://github.com/ranaroussi/yfinance",
        "scikit-learn Ridge Regression documentation. https://scikit-learn.org/stable/modules/generated/sklearn.linear_model.Ridge.html",
        "Hyndman, R. J., and Athanasopoulos, G. Forecasting Principles and Practice. https://otexts.com/fpp3/tscv.html",
        "Project repository. https://github.com/Siddhantkr07/ML-Minor",
    )
    for reference in references:
        add_bullet(document, reference)

    document.save(REPORT_PATH)
    print(f"Created {REPORT_PATH}")


if __name__ == "__main__":
    main()
